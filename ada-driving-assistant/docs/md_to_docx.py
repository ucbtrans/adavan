#!/usr/bin/env python
"""Convert a Markdown doc in this folder to a styled .docx.

Handles the subset of Markdown these docs actually use: h1-h3, paragraphs,
bullet lists, pipe tables, horizontal rules, and inline **bold**, *italic*,
`code` and [links](target).

    python docs/md_to_docx.py docs/FLEET_MANAGER_GUIDE.md
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"
INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
CODE_INK = RGBColor(0xA3, 0x1D, 0x1D)

# **bold** | *italic* | `code` | [text](href)
INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_runs(par, text):
    """Write text into a paragraph, honouring inline Markdown."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = par.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = par.add_run(chunk[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(9.5)
            run.font.color.rgb = CODE_INK
        elif chunk.startswith("[") and "](" in chunk:
            label, href = re.match(r"\[([^\]]+)\]\(([^)]+)\)", chunk).groups()
            label = label.strip("`*")          # link text may itself be `code`
            run = par.add_run(label)
            run.font.color.rgb = ACCENT
            run.underline = True
            if not href.startswith("#"):
                par.add_run(f" ({href})").font.size = Pt(8.5)
        elif chunk.startswith("*") and chunk.endswith("*"):
            run = par.add_run(chunk[1:-1])
            run.italic = True
        else:
            par.add_run(chunk)


def shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def hrule(doc):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(10)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "C8C8C8")
    pbdr.append(bottom)
    par._p.get_or_add_pPr().append(pbdr)


def build(md_path, docx_path):
    lines = open(md_path, encoding="utf-8").read().split("\n")

    doc = Document()
    for name, size in (("Normal", 11), ("List Bullet", 11)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.color.rgb = INK
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after = Pt(8)
    normal.line_spacing = 1.15

    for lvl, size, color in ((1, 22, ACCENT), (2, 15, ACCENT), (3, 12, INK)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True

    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(0.9)
        sec.left_margin = sec.right_margin = Inches(1.0)

    i, n = 0, len(lines)
    while i < n:
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            hrule(doc)
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            par = doc.add_heading(level=level)
            par.paragraph_format.space_before = Pt(16 if level > 1 else 0)
            par.paragraph_format.space_after = Pt(6)
            add_runs(par, m.group(2))
            for run in par.runs:                      # heading style owns the colour
                run.font.color.rgb = ACCENT if level < 3 else INK
                run.font.name = BODY_FONT
            i += 1
            continue

        # pipe table: header row, separator, body rows
        if stripped.startswith("|") and i + 1 < n and re.match(
                r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            def cells(row):
                return [c.strip() for c in row.strip().strip("|").split("|")]

            header = cells(lines[i])
            body = []
            j = i + 2
            while j < n and lines[j].strip().startswith("|"):
                body.append(cells(lines[j]))
                j += 1

            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, text in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ""
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(2)
                add_runs(par, text)
                for run in par.runs:
                    run.bold = True
                shade(cell, "E8EEF5")
            for row in body:
                cells_out = table.add_row().cells
                for c, text in enumerate(row[:len(header)]):
                    cells_out[c].text = ""
                    par = cells_out[c].paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    add_runs(par, text)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i = j
            continue

        if re.match(r"^[-*]\s+", stripped):
            # a bullet may wrap onto following indented lines
            text = re.sub(r"^[-*]\s+", "", stripped)
            i += 1
            while i < n and lines[i].startswith("  ") and lines[i].strip() \
                    and not re.match(r"^\s*[-*]\s+", lines[i]):
                text += " " + lines[i].strip()
                i += 1
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(4)
            add_runs(par, text)
            continue

        # paragraph: join wrapped lines until a blank or a block starter
        buf = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith("#") or nxt.startswith("|")
                    or nxt in ("---", "***", "___") or re.match(r"^[-*]\s+", nxt)):
                break
            buf.append(nxt)
            i += 1
        text = " ".join(buf)

        par = doc.add_paragraph()
        # a standalone *(diagram block …)* line reads as a caption
        if text.startswith("*(") and text.endswith(")*"):
            par.paragraph_format.space_after = Pt(10)
            run = par.add_run(text[1:-1])
            run.italic = True
            run.font.color.rgb = MUTED
            run.font.size = Pt(10)
        else:
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(par, text)

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "docs/FLEET_MANAGER_GUIDE.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + ".docx"
    out = build(src, dst)
    print(f"wrote {out} ({os.path.getsize(out)/1024:.0f} KB)")
